#!/usr/bin/env python3
"""Build src/data/comuni.json from ISTAT sources.

Downloads:
- Elenco dei comuni italiani (21 February 2026 permalink)
- Confini delle unità amministrative a fini statistici, 1 January 2026,
  generalised

The boundary file is the 1 January 2026 release (7,896 comuni). The code list
is the 21 February 2026 release (7,894 comuni). This script follows the code
list. Castegnero Nanto (024129) is the union of Castegnero (024027) and Nanto
(024071). Lirio (018082) is in the January boundaries but not in the February
code list, so it is left out.

Requires: pyshp, shapely, pyproj, openpyxl
"""

from __future__ import annotations

import json
import sys
import tempfile
import urllib.request
import zipfile
from pathlib import Path

import openpyxl
import shapefile
from pyproj import Transformer
from shapely.geometry import shape as to_shape
from shapely.ops import unary_union

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "src" / "data" / "comuni.json"

CODES_URL = "https://www.istat.it/storage/codici-unita-amministrative/Elenco-comuni-italiani.xlsx"
BOUNDS_URL = "https://www.istat.it/storage/cartografia/confini_amministrativi/generalizzati/2026/Limiti01012026_g.zip"

# New comune instituted 21 February 2026, after the 1 January boundary release.
MERGED = {
    "024129": ("024027", "024071"),
}

USER_AGENT = "GuidocasaComuniBuild/1.0 (dataset regeneration; +https://www.istat.it)"


def download(url: str, dest: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=120) as response:
        dest.write_bytes(response.read())


def load_codes(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    next(rows)
    comuni = []
    for row in rows:
        if not row or row[4] is None:
            continue
        code = str(row[4]).strip().zfill(6)
        name = str(row[6]).strip()
        alt = str(row[7]).strip() if row[7] else ""
        province = str(row[11]).strip()
        sigla = str(row[14]).strip()
        comuni.append(
            {
                "istat": code,
                "name": name,
                "alt": alt,
                "province": province,
                "sigla": sigla,
            }
        )
    return comuni


def geometry_by_code(shp_path: Path) -> dict:
    reader = shapefile.Reader(str(shp_path), encoding="utf-8")
    found = {}
    for index, record in enumerate(reader.iterRecords()):
        data = record.as_dict()
        code = str(data["PRO_COM_T"]).strip()
        geom = to_shape(reader.shape(index).__geo_interface__)
        area_m2 = float(data["Shape_Area"])
        if code in found:
            found[code] = (unary_union([found[code][0], geom]), found[code][1] + area_m2)
        else:
            found[code] = (geom, area_m2)
    return found


def point_inside(geom):
    centroid = geom.centroid
    if geom.contains(centroid):
        return centroid
    return geom.representative_point()


def main() -> None:
    transformer = Transformer.from_crs("EPSG:32632", "EPSG:4326", always_xy=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        xlsx = tmp_path / "elenco.xlsx"
        archive = tmp_path / "limiti.zip"
        print("Downloading code list…", file=sys.stderr)
        download(CODES_URL, xlsx)
        print("Downloading boundaries…", file=sys.stderr)
        download(BOUNDS_URL, archive)
        with zipfile.ZipFile(archive) as zf:
            zf.extractall(tmp_path)
        shp = next(tmp_path.glob("Com*_g/*_WGS84.shp"))
        geometries = geometry_by_code(shp)
        codes = load_codes(xlsx)

    records = []
    missing = []
    for comune in codes:
        code = comune["istat"]
        if code in MERGED:
            parts = []
            area = 0.0
            for part in MERGED[code]:
                if part not in geometries:
                    missing.append(code)
                    break
                geom, part_area = geometries[part]
                parts.append(geom)
                area += part_area
            else:
                geom = unary_union(parts)
                point = point_inside(geom)
                lng, lat = transformer.transform(point.x, point.y)
                records.append(row(comune, lat, lng, area))
            continue
        if code not in geometries:
            missing.append(code)
            continue
        geom, area = geometries[code]
        point = point_inside(geom)
        lng, lat = transformer.transform(point.x, point.y)
        records.append(row(comune, lat, lng, area))

    if missing:
        raise SystemExit(f"Missing geometry for: {', '.join(missing)}")

    records.sort(key=lambda item: item[0])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    pretty = ",\n".join(json.dumps(item, ensure_ascii=False) for item in records)
    OUT.write_text("[\n" + pretty + "\n]\n", encoding="utf-8")
    print(f"Wrote {len(records)} comuni to {OUT}", file=sys.stderr)


def row(comune: dict, lat: float, lng: float, area_m2: float) -> list:
    return [
        comune["istat"],
        comune["name"],
        comune["alt"],
        comune["province"],
        comune["sigla"],
        round(lat, 5),
        round(lng, 5),
        round(area_m2 / 1_000_000, 2),
    ]


if __name__ == "__main__":
    main()
